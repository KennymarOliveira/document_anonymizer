from abc import ABC, abstractmethod
import io
import os
from typing import Tuple

import docx

from app.core.configuration.settings import settings
from app.core.security.logger import get_logger

logger = get_logger(__name__)

GOOGLE_DOCS_MIMETYPE = "application/vnd.google-apps.document"
DOCX_MIMETYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


class GoogleDriveClientInterface(ABC):
    """Interface abstrata para clientes de integração com o Google Drive."""

    @abstractmethod
    def fetch_file(self, file_id: str) -> Tuple[str, bytes]:
        """Recupera o arquivo do Google Drive por ID.

        Retorna:
            Tuple[str, bytes]: (nome_do_arquivo_com_extensao, conteudo_em_bytes)
        """
        pass


class MockGoogleDriveClient(GoogleDriveClientInterface):
    """Cliente Mock para desenvolvimento local sem credenciais de Service Account."""

    def __init__(self) -> None:
        self._mock_database = {
            "doc-corporativo-123": {
                "name": "Contrato_Prestacao_Servicos.docx",
                "content": self._generate_sample_docx(),
            },
            "doc-juridico-456": {
                "name": "Peticao_Inicial.txt",
                "content": (
                    "EXCELENTÍSSIMO SENHOR DOUTOR JUIZ DE DIREITO DA 1ª VARA CÍVEL.\n"
                    "Autor: João da Silva, brasileiro, CPF nº 123.456.789-00, residente em São Paulo.\n"
                    "Réu: Banco Exemplo S.A., CNPJ nº 00.123.456/0001-99."
                ).encode("utf-8"),
            },
        }

    def _generate_sample_docx(self) -> bytes:
        doc = docx.Document()
        doc.add_heading("Contrato de Prestação de Serviços Jurídicos", level=1)
        doc.add_paragraph(
            "Contratante: Roberto Almeida, CPF 987.654.321-11, residente na Av. Paulista, 1000."
        )
        doc.add_paragraph(
            "Contratada: Advocacia Silva & Santos, CNPJ 12.345.678/0001-90, com sede em Curitiba."
        )
        stream = io.BytesIO()
        doc.save(stream)
        stream.seek(0)
        return stream.getvalue()

    def fetch_file(self, file_id: str) -> Tuple[str, bytes]:
        logger.info("[MOCK] Buscando arquivo com file_id: %s", file_id)
        if file_id not in self._mock_database:
            raise FileNotFoundError(
                f"Arquivo com ID '{file_id}' não encontrado no Google Drive (MOCK)."
            )

        item = self._mock_database[file_id]
        return item["name"], item["content"]


class GoogleDriveClient(GoogleDriveClientInterface):
    """Cliente real usando Service Account da Google Drive API.

    Esta classe depende de credenciais reais do Google Cloud Console.
    """

    def __init__(self, credentials_path: str | None = None) -> None:
        self.credentials_path = credentials_path or os.getenv("GOOGLE_APPLICATION_CREDENTIALS")
        self._service = None

    def _get_service(self):
        """Inicializa o serviço do Google Drive caso as credenciais estejam disponíveis."""
        if self._service is not None:
            return self._service

        if not self.credentials_path or not os.path.exists(self.credentials_path):
            raise RuntimeError(
                f"Credenciais de Service Account não encontradas em: {self.credentials_path}. "
                "Defina a variável GOOGLE_APPLICATION_CREDENTIALS ou utilize o MockGoogleDriveClient."
            )

        # Import lazy para não exigir a lib caso o ambiente ainda não a tenha instalada
        from google.oauth2 import service_account
        from googleapiclient.discovery import build

        scopes = ["https://www.googleapis.com/auth/drive.readonly"]
        creds = service_account.Credentials.from_service_account_file(
            self.credentials_path, scopes=scopes
        )
        self._service = build("drive", "v3", credentials=creds)
        return self._service

    def fetch_file(self, file_id: str) -> Tuple[str, bytes]:
        from googleapiclient.errors import HttpError
        from googleapiclient.http import MediaIoBaseDownload

        logger.info("Buscando arquivo com file_id: %s", file_id)
        service = self._get_service()
        request = None
        stream = io.BytesIO()

        try:
            metadata = service.files().get(fileId=file_id, fields="name, mimeType").execute()
        except HttpError as exc:
            if exc.resp.status == 404:
                raise FileNotFoundError(f"Arquivo '{file_id}' não encontrado no Google Drive.") from exc
            raise

        file_name = metadata.get("name", file_id)
        file_type = metadata.get("mimeType", "")
        

        if file_type == GOOGLE_DOCS_MIMETYPE:
            request = service.files().export_media(fileId=file_id, mimeType=DOCX_MIMETYPE)
            if not file_name.lower().endswith(".docx"):
                file_name = f"{file_name}.docx"
        else:
            request = service.files().get_media(fileId=file_id)

        downloader = MediaIoBaseDownload(stream, request)
        done = False
        while not done:
            status, done = downloader.next_chunk()
        content_bytes = stream.getvalue()

        return file_name, content_bytes 


def get_google_drive_client() -> GoogleDriveClientInterface:
    """Factory com fallback automático: se houver credenciais reais, usa o cliente real;
    caso contrário, usa o Mock.
    """
    creds_path = settings.GOOGLE_APPLICATION_CREDENTIALS
    use_mock = settings.USE_GOOGLE_DRIVE_MOCK

    if not use_mock and creds_path and os.path.exists(creds_path):
        logger.info("Utilizando GoogleDriveClient real com credenciais de Service Account.")
        return GoogleDriveClient(credentials_path=creds_path)

    logger.info("Utilizando MockGoogleDriveClient (ambiente de testes/sem credenciais).")
    return MockGoogleDriveClient()
