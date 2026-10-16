from collections import defaultdict
import io
from typing import Any, Dict, List, Literal

import docx
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import RGBColor
import pymupdf

REDACTION_MODE_BLACKOUT = "blackout"
REDACTION_MODE_BLACK_WHITE_TEXT = "black_white_text"

RedactionMode = Literal[
    "blackout",
    "black_white_text",
    "tarja_preta",
    "tarja_texto_branco",
]


def normalize_redaction_mode(mode: str | None) -> str:
    """Normaliza o modo de tarja para os tipos canônicos."""
    if not mode:
        return REDACTION_MODE_BLACKOUT
    m = mode.strip().lower()
    blackout_aliases = [
        "blackout",
        "tarja_preta",
        "solid",
        "black",
        "preto",
        "padrao",
        "default",
    ]
    if m in blackout_aliases:
        return REDACTION_MODE_BLACKOUT
    white_text_aliases = [
        "black_white_text",
        "texto_branco",
        "tarja_texto_branco",
        "black_bg_white_text",
        "white_on_black",
        "white_text",
        "tarja_com_texto",
    ]
    if m in white_text_aliases:
        return REDACTION_MODE_BLACK_WHITE_TEXT
    raise ValueError(
        f"Modo de tarja '{mode}' não suportado. Opções válidas: "
        "'blackout' (tarja preta sólida) ou "
        "'black_white_text' (tarja preta de fundo com texto branco)."
    )


def build_txt_entity_identifiers(
    entities: List[Dict[str, Any]],
) -> Dict[tuple[str, str], str]:
    """Cria IDs estáveis por tipo e valor para a versão TXT de testes."""
    identifiers: Dict[tuple[str, str], str] = {}
    label_counts: Dict[str, int] = defaultdict(int)

    for entity in entities:
        value = entity.get("text")
        if not value:
            continue
        label = str(entity.get("label") or "DADO").upper()
        key = (label, value)
        if key not in identifiers:
            label_counts[label] += 1
            identifiers[key] = f"[{label}_{label_counts[label]:03d}]"

    return identifiers


def _style_run_black_white_text(run: Any) -> None:
    """Aplica formatação de texto branco com fundo preto em um run."""
    run.font.color.rgb = RGBColor(255, 255, 255)
    ns = nsdecls("w")
    shd = parse_xml(f'<w:shd {ns} w:fill="000000"/>')
    run._r.get_or_add_rPr().append(shd)


def _is_run_styled_white_on_black(run: Any) -> bool:
    """Verifica se o run já está com a cor branca configurada."""
    return (
        run.font.color is not None
        and getattr(run.font.color, "rgb", None) == RGBColor(255, 255, 255)
    )


def _apply_black_white_to_paragraph(paragraph: Any, targets: List[str]) -> None:
    """Divide os runs do parágrafo aplicando fundo preto e texto branco."""
    for target in targets:
        if not target or target not in paragraph.text:
            continue

        # Consolida runs se o termo estiver quebrado entre runs sem estilo
        has_unstyled = any(
            target in r.text
            for r in paragraph.runs
            if not _is_run_styled_white_on_black(r)
        )
        if not has_unstyled:
            styled_runs = [
                r for r in paragraph.runs if _is_run_styled_white_on_black(r)
            ]
            if not styled_runs:
                paragraph.text = paragraph.text

        runs = list(paragraph.runs)
        for r in runs:
            if _is_run_styled_white_on_black(r):
                continue
            if target in r.text:
                parts = r.text.split(target)
                r.text = parts[0]
                cur_r = r
                for i in range(1, len(parts)):
                    new_run = paragraph.add_run(target)
                    _style_run_black_white_text(new_run)
                    cur_r._r.addnext(new_run._r)
                    cur_r = new_run
                    if parts[i]:
                        after_run = paragraph.add_run(parts[i])
                        cur_r._r.addnext(after_run._r)
                        cur_r = after_run

                if not r.text and r._r.getparent() is not None:
                    r._r.getparent().remove(r._r)


def anonymize_docx_in_place(
    content: bytes,
    entities: List[Dict[str, Any]],
    redaction_mode: str = "blackout",
) -> io.BytesIO:
    """Anonimiza um DOCX com tarjas sólidas ou texto em contraste."""
    mode = normalize_redaction_mode(redaction_mode)
    doc = docx.Document(io.BytesIO(content))

    if mode == REDACTION_MODE_BLACKOUT:
        replacements = {
            e["text"]: "█" * len(e["text"])
            for e in entities
            if e.get("text")
        }

        def _replace_in_paragraph(paragraph):
            for old_text, new_text in replacements.items():
                if old_text in paragraph.text:
                    replaced = False
                    for run in paragraph.runs:
                        if old_text in run.text:
                            run.text = run.text.replace(old_text, new_text)
                            replaced = True

                    if not replaced:
                        paragraph.text = paragraph.text.replace(
                            old_text, new_text
                        )

        process_p = _replace_in_paragraph
    else:
        targets = sorted(
            {
                e["text"]
                for e in entities
                if e.get("text") and e["text"].strip()
            },
            key=len,
            reverse=True,
        )

        def _replace_in_paragraph(paragraph):
            _apply_black_white_to_paragraph(paragraph, targets)

        process_p = _replace_in_paragraph

    for p in doc.paragraphs:
        process_p(p)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    process_p(p)

    for section in doc.sections:
        for p in section.header.paragraphs:
            process_p(p)
        for p in section.footer.paragraphs:
            process_p(p)

    out_stream = io.BytesIO()
    doc.save(out_stream)
    out_stream.seek(0)
    return out_stream


def anonymize_pdf_in_place(
    content: bytes,
    entities: List[Dict[str, Any]],
    redaction_mode: str = "blackout",
) -> io.BytesIO:
    """Anonimiza um PDF com faixas pretas ou fundo preto + texto branco."""
    mode = normalize_redaction_mode(redaction_mode)
    pdf_doc = pymupdf.open(stream=content, filetype="pdf")

    page_terms_map: Dict[int, set[str]] = defaultdict(set)
    global_terms: set[str] = set()
    page_rects_map: Dict[
        int, list[tuple[pymupdf.Rect, str, str, int]]
    ] = defaultdict(list)
    global_rects: list[tuple[pymupdf.Rect, str, str, int]] = []

    for e in entities:
        text = e.get("text", "").strip()
        if not text:
            continue
        page_num = e.get("page")
        tier = int(e.get("tier", 50))
        entity_rects = e.get("rects")
        if entity_rects:
            targets = (
                page_rects_map[page_num]
                if page_num is not None
                else global_rects
            )
            render_segments = e.get("render_segments") or []
            if render_segments:
                for segment in render_segments:
                    segment_rect = segment.get("rect")
                    segment_text = str(segment.get("text") or text)
                    if segment_rect:
                        targets.append(
                            (pymupdf.Rect(segment_rect), text, segment_text, tier)
                        )
            else:
                for raw_rect in entity_rects:
                    targets.append((pymupdf.Rect(raw_rect), text, text, tier))
            continue
        if page_num is not None:
            page_terms_map[page_num].add(text)
        else:
            global_terms.add(text)

    def _draw_white_text(page, rect: pymupdf.Rect, text: str) -> None:
        """Desenha texto conhecido sem reextrair conteúdo do retângulo do PDF.

        Se não houver espaço mesmo com fonte reduzida, não desenha nada. Assim a
        tarja nunca avança para texto vizinho nem produz uma substituição parcial.
        """
        # Primeiro preserva uma pequena margem; se necessário, tenta novamente
        # com toda a área da tarja e reduz a fonte até 0,5 pt.
        for padding in (0.5, 0):
            draw_rect = pymupdf.Rect(rect)
            draw_rect.x0 += padding
            draw_rect.y0 += padding
            draw_rect.x1 -= padding
            draw_rect.y1 -= padding
            if draw_rect.is_empty or draw_rect.is_infinite:
                continue

            font_size = min(10.0, draw_rect.height * 0.75)
            while font_size >= 0.5:
                remaining = page.insert_textbox(
                    draw_rect,
                    text,
                    fontname="helv",
                    fontsize=font_size,
                    color=(1, 1, 1),
                    align=pymupdf.TEXT_ALIGN_CENTER,
                    overlay=True,
                )
                if remaining >= 0:
                    return
                font_size -= 0.5

    def _redact_page(page_idx: int) -> None:
        page = pdf_doc[page_idx]
        page_num = page_idx + 1
        terms_for_page = page_terms_map.get(page_num, set()) | global_terms
        positioned_candidates = page_rects_map.get(page_num, []) + global_rects
        if not terms_for_page and not positioned_candidates:
            return

        candidates: List[tuple[pymupdf.Rect, str, str, int]] = list(positioned_candidates)

        single_word_terms = {t for t in terms_for_page if " " not in t}
        multi_word_terms = sorted(
            [t for t in terms_for_page if " " in t], key=len, reverse=True
        )

        if single_word_terms:
            single_terms_lower = {t.lower(): t for t in single_word_terms}

            for w in page.get_text("words"):
                raw_l = w[4].lower()
                clean_l = raw_l.strip(".,;:()[]{}\\'\"")
                matched_term = None
                if raw_l in single_terms_lower:
                    matched_term = single_terms_lower[raw_l]
                elif clean_l in single_terms_lower:
                    matched_term = single_terms_lower[clean_l]

                if matched_term:
                    rect = pymupdf.Rect(w[:4])
                    candidates.append((rect, matched_term, matched_term, 0))

        # A busca textual do PDF aceita correspondências parciais. Por isso,
        # termos de uma palavra que não foram encontrados como palavras inteiras
        # não podem usar search_for: "PR" não deve atingir "PRESTAÇÃO".
        for term in multi_word_terms:
            for rect in page.search_for(term):
                candidates.append((rect, term, term, 0))

        # Prefere a expressão mais longa e evita anotações visuais sobrepostas.
        candidates.sort(key=lambda item: (-item[3], -len(item[1])))
        selected: List[tuple[pymupdf.Rect, str]] = []
        for rect, term, draw_text, _ in candidates:
            if any(rect.intersects(selected_rect) for selected_rect, _ in selected):
                continue
            selected.append((rect, draw_text))

        for rect, _ in selected:
            page.add_redact_annot(rect, fill=(0, 0, 0))

        page.apply_redactions()
        if mode == REDACTION_MODE_BLACK_WHITE_TEXT:
            for rect, text in selected:
                _draw_white_text(page, rect, text)

    num_pages = len(pdf_doc)
    for idx in range(num_pages):
        _redact_page(idx)

    out_stream = io.BytesIO()
    pdf_doc.save(out_stream, garbage=1, deflate=True)
    out_stream.seek(0)
    return out_stream


def build_anonymized_file(
    filename: str,
    original_content: bytes,
    entities: List[Dict[str, Any]],
    redaction_mode: str = "blackout",
) -> tuple[io.BytesIO, str]:
    """Direciona o arquivo para o construtor aplicando o modo de tarja."""
    mode = normalize_redaction_mode(redaction_mode)
    ext = filename.split(".")[-1].lower()

    if ext == "txt":
        text = original_content.decode("utf-8")
        entities_sorted = sorted(
            entities, key=lambda x: len(x.get("text", "")), reverse=True
        )

        if mode == REDACTION_MODE_BLACK_WHITE_TEXT:
            # TXT não suporta formatação visual portável. Use identificadores
            # rastreáveis para teste sem manter o valor sensível no arquivo.
            replacements = build_txt_entity_identifiers(entities)

            for entity in entities_sorted:
                value = entity.get("text")
                if value:
                    label = str(entity.get("label") or "DADO").upper()
                    text = text.replace(value, replacements[(label, value)])
        else:
            for entity in entities_sorted:
                value = entity.get("text")
                if value:
                    text = text.replace(value, "█" * len(value))
        return io.BytesIO(text.encode("utf-8")), "text/plain"

    if ext in ["doc", "docx"]:
        mime_type = (
            "application/vnd.openxmlformats-officedocument."
            "wordprocessingml.document"
        )
        return (
            anonymize_docx_in_place(
                original_content, entities, redaction_mode=mode
            ),
            mime_type,
        )

    if ext == "pdf":
        return (
            anonymize_pdf_in_place(
                original_content, entities, redaction_mode=mode
            ),
            "application/pdf",
        )

    raise ValueError(f"Extensão não suportada para reconstrução: {ext}")
    raise ValueError(f"Extensão não suportada para reconstrução: {ext}")
